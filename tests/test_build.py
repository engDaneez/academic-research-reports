from pathlib import Path
from tempfile import TemporaryDirectory
import json
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(ROOT / 'scripts'))
from build_docx import build, _validate


class TestDocxGenerator(unittest.TestCase):
    def test_example_report(self):
        data = json.loads((ROOT / 'assets/example-report.json').read_text(encoding='utf-8'))
        with TemporaryDirectory() as td:
            out = Path(td) / 'report.docx'
            build(data, out)
            self.assertGreater(out.stat().st_size, 10000)
            d = Document(out)
            self.assertEqual(len(d.sections), 2)
            self.assertTrue(any(p.style.name == 'Heading 2' for p in d.paragraphs))
            self.assertTrue(any(p.text.strip() == 'المقدمة' for p in d.paragraphs))

    def test_research_front_matter(self):
        data = json.loads((ROOT / 'assets/example-graduation-research.json').read_text(encoding='utf-8'))
        with TemporaryDirectory() as td:
            out = Path(td) / 'thesis.docx'
            build(data, out)
            d = Document(out)
            self.assertTrue(any(p.text.strip() == 'الإهداء' for p in d.paragraphs))

    def test_missing_data(self):
        with self.assertRaises(ValueError):
            _validate({'title': 'x', 'language': 'ar', 'sections': []})


if __name__ == '__main__':
    unittest.main()
