#!/usr/bin/env python3
"""Basic structural audit; does not replace full visual layout inspection."""
import argparse
from pathlib import Path
from docx import Document


def main():
    p = argparse.ArgumentParser()
    p.add_argument("docx", type=Path)
    a = p.parse_args()
    doc = Document(a.docx)
    headings = [(p.style.name, p.text.strip()) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    texts = [p.text for p in doc.paragraphs]
    missing = []
    if not headings:
        missing.append("No real Heading styles")
    if len(doc.sections) < 2:
        missing.append("No distinct front/main sections")
    if not any(x.strip() for x in texts):
        missing.append("Document has no text")
    print(f"Paragraphs: {len(doc.paragraphs)} | Headings: {len(headings)} | Sections: {len(doc.sections)}")
    print("First headings:", ", ".join(title for _, title in headings[:8]))
    if missing:
        for issue in missing:
            print("ISSUE:", issue)
        raise SystemExit(1)
    print("Structural check passed. Visual page-by-page inspection is still REQUIRED.")


if __name__ == "__main__":
    main()
