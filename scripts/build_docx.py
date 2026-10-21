#!/usr/bin/env python3
"""Build a modest, styled DOCX from a UTF-8 JSON document.

This is a fallback formatter, not a research writer or citation verifier.
A generated Word TOC may require updating in Microsoft Word.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def _set_rtl(paragraph: Any) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1")


def _rtl_run(run: Any) -> None:
    rpr = run._r.get_or_add_rPr()
    rtl = rpr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        rpr.append(rtl)
    rtl.set(qn("w:val"), "1")


def _set_style_font(style: Any, name: str, size: int, bold: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn("w:" + attr), name)


def _set_text(paragraph: Any, value: str, arabic: bool) -> None:
    run = paragraph.add_run(value)
    if arabic:
        _rtl_run(run)


def _apply_page(page: Any, margins: dict, right_cm: float | None = None) -> None:
    page.page_height = Cm(29.7)
    page.page_width = Cm(21)
    page.top_margin = Cm(float(margins.get("top_cm", 2.5)))
    page.bottom_margin = Cm(float(margins.get("bottom_cm", 2.5)))
    page.left_margin = Cm(float(margins.get("left_cm", 2.5)))
    page.right_margin = Cm(float(right_cm if right_cm is not None else margins.get("right_cm", 2.5)))


def _toc_field(paragraph: Any, arabic: bool) -> None:
    run = paragraph.add_run()
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(start)
    run._r.append(instr)
    run._r.append(sep)
    _set_text(paragraph, "افتح الملف في Word ثم حدّث الفهرس." if arabic else "Update this TOC field in Word.", arabic)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def _page_number(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def _restart_numbering(section: Any) -> None:
    sectpr = section._sectPr
    pg_num = sectpr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sectpr.append(pg_num)
    pg_num.set(qn("w:start"), "1")


def _configure_styles(doc: Document, arabic: bool, options: dict) -> None:
    font = options.get("font_name", "Traditional Arabic" if arabic else "Times New Roman")
    body = int(options.get("body_font_size", 14 if arabic else 12))
    h1 = int(options.get("heading_1_size", 18 if arabic else 16))
    h2 = int(options.get("heading_2_size", 16 if arabic else 14))
    for name, size, bold in [("Normal", body, False), ("Title", 22 if arabic else 20, True),
                             ("Heading 1", h1, True), ("Heading 2", h2, True),
                             ("Heading 3", max(body, h2 - 2), True)]:
        _set_style_font(doc.styles[name], font, size, bold)
        fmt = doc.styles[name].paragraph_format
        fmt.space_after = Pt(float(options.get("space_after_pt", 6)))
        fmt.line_spacing = float(options.get("line_spacing", 1.5))
        if arabic:
            ppr = doc.styles[name]._element.get_or_add_pPr()
            bidi = ppr.find(qn("w:bidi"))
            if bidi is None:
                bidi = OxmlElement("w:bidi")
                ppr.append(bidi)
            bidi.set(qn("w:val"), "1")
    # Built-in Word Title sometimes has a decorative bottom border; a formal
    # academic layout defaults to clean black headings without this border.
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)


def _paragraph(doc: Document, text: str, arabic: bool, style: str | None = None, center: bool = False) -> Any:
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else (
        WD_ALIGN_PARAGRAPH.JUSTIFY if style is None else
        (WD_ALIGN_PARAGRAPH.RIGHT if arabic else WD_ALIGN_PARAGRAPH.LEFT)
    )
    if arabic:
        _set_rtl(p)
    _set_text(p, text, arabic)
    return p


def _validate(content: dict) -> None:
    if not isinstance(content, dict) or not isinstance(content.get("title"), str) or not content["title"].strip():
        raise ValueError("Input must contain a non-empty 'title' string")
    if content.get("language", "ar") not in ("ar", "en"):
        raise ValueError("'language' must be 'ar' or 'en'")
    sections = content.get("sections")
    if not isinstance(sections, list) or not sections:
        raise ValueError("'sections' must be a non-empty list")
    for i, section in enumerate(sections, 1):
        if not isinstance(section, dict) or not isinstance(section.get("title"), str):
            raise ValueError(f"Section {i}: missing title")
        if section.get("level", 1) not in (1, 2, 3):
            raise ValueError(f"Section {i}: level must be 1, 2 or 3")
        if not isinstance(section.get("paragraphs", []), list):
            raise ValueError(f"Section {i}: paragraphs must be an array")
    if not isinstance(content.get("references", []), list):
        raise ValueError("'references' must be an array")


def build(content: dict, output: Path) -> None:
    _validate(content)
    arabic = content.get("language", "ar") == "ar"
    options = content.get("format", {})
    doc = Document()
    _configure_styles(doc, arabic, options)
    _apply_page(doc.sections[0], options)

    cover = content.get("cover", {})
    for key in ("institution", "college", "department"):
        if cover.get(key):
            _paragraph(doc, str(cover[key]), arabic, center=True)
    doc.add_paragraph()
    _paragraph(doc, content["title"], arabic, style="Title", center=True)
    doc.add_paragraph()
    for key, arlabel, enlabel in [
        ("prepared_by", "إعداد", "Prepared by"),
        ("supervisor", "إشراف", "Supervised by"),
        ("academic_year", "السنة الدراسية", "Academic year"),
    ]:
        if cover.get(key):
            _paragraph(doc, f"{arlabel if arabic else enlabel}: {cover[key]}", arabic, center=True)

    front = content.get("front_matter", [])
    include_toc = content.get("toc", False)
    if front or include_toc:
        doc.add_page_break()

    for index, item in enumerate(front):
        if not isinstance(item, dict):
            raise ValueError(f"Front matter item {index+1} must be an object")
        _paragraph(doc, str(item.get("title", "")), arabic, style="Heading 1", center=True)
        for itemtext in item.get("paragraphs", []):
            _paragraph(doc, str(itemtext), arabic)
        if index != len(front)-1 or include_toc:
            doc.add_page_break()

    if include_toc:
        _paragraph(doc, "فهرس المحتويات" if arabic else "Table of Contents", arabic, style="Heading 1", center=True)
        p = doc.add_paragraph()
        if arabic:
            _set_rtl(p)
        _toc_field(p, arabic)

    main_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    _apply_page(main_section, options)
    main_section.footer.is_linked_to_previous = False
    _restart_numbering(main_section)
    footer = main_section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT if arabic else WD_ALIGN_PARAGRAPH.CENTER
    if arabic:
        # Override the inherited bidi Normal style for bottom-left page numbers.
        ppr = footer._p.get_or_add_pPr()
        bidi = ppr.find(qn("w:bidi"))
        if bidi is None:
            bidi = OxmlElement("w:bidi")
            ppr.append(bidi)
        bidi.set(qn("w:val"), "0")
    _page_number(footer)

    for i, sec in enumerate(content["sections"]):
        level = sec.get("level", 1)
        _paragraph(doc, str(sec["title"]), arabic, style=f"Heading {level}")
        for para in sec.get("paragraphs", []):
            if not isinstance(para, str):
                raise ValueError(f"Section {i+1} paragraphs must be text")
            _paragraph(doc, para, arabic)

    references = content.get("references", [])
    if references:
        _paragraph(doc, "المصادر والمراجع" if arabic else "References", arabic, style="Heading 1")
        for ref in references:
            reftext = ref if isinstance(ref, str) else ref.get("text", "")
            if not reftext:
                raise ValueError("Each reference must provide text")
            p = _paragraph(doc, str(reftext), arabic)
            p.paragraph_format.first_line_indent = Cm(-0.6) if not arabic else Cm(0)
            p.paragraph_format.left_indent = Cm(0.6) if not arabic else Cm(0)

    # Ask Word to update fields when it opens the document.
    settings = doc.settings._element
    upd = settings.find(qn("w:updateFields"))
    if upd is None:
        upd = OxmlElement("w:updateFields")
        settings.append(upd)
    upd.set(qn("w:val"), "true")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    print(f"Created: {output} (content_sections={len(content['sections'])}; front_matter={len(front)}; references={len(references)})")
    if include_toc:
        print("Reminder: update the table of contents in Microsoft Word and visually review all pages.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True, help="Input UTF-8 JSON file")
    parser.add_argument("--output", type=Path, required=True, help="DOCX output path")
    args = parser.parse_args()
    with args.input.open("r", encoding="utf-8") as f:
        content = json.load(f)
    build(content, args.output)


if __name__ == "__main__":
    main()
